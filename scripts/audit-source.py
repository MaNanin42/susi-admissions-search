import argparse
import collections
import hashlib
import json
from pathlib import Path

import openpyxl
from pypdf import PdfReader


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open('rb') as stream:
        for chunk in iter(lambda: stream.read(1024 * 1024), b''):
            digest.update(chunk)
    return digest.hexdigest().upper()


def clean(value) -> str:
    return '' if value is None else str(value).strip()


def value_text(value):
    text = clean(value)
    if not text or text == '-':
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return text


def source_record(row):
    record = {
        'r': value_text(row[0]),
        'u': value_text(row[3]),
        'c': value_text(row[8]),
        'm': value_text(row[7]),
    }
    for prefix, columns in [
        ('k1', (11, 12, 13)), ('k2', (17, 18, 19)), ('k3', (23, 24, 25)),
        ('j1', (29, 30, 31)), ('j2', (35, 36, 37)),
    ]:
        for suffix, column in zip(('n', 'a', 'b'), columns):
            value = value_text(row[column])
            if value is not None:
                record[prefix + suffix] = value
    return {key: value for key, value in record.items() if value is not None}


def baseline_core(record):
    corrections = {'k1navi50', 'k1navi70', 'k1appAvg', 'k1metric'}
    return {key: value for key, value in record.items() if key not in corrections}


def signature(record):
    return json.dumps(record, ensure_ascii=False, sort_keys=True, separators=(',', ':'))


def main() -> None:
    parser = argparse.ArgumentParser(description='Audit official NAVI sources against the V10.6 public baseline.')
    parser.add_argument('--workbook', required=True)
    parser.add_argument('--guidebook', required=True)
    parser.add_argument('--output')
    args = parser.parse_args()

    root = Path(__file__).resolve().parents[1]
    workbook_path = Path(args.workbook).resolve()
    guidebook_path = Path(args.guidebook).resolve()
    baseline = json.loads((root / 'data' / 'baseline' / 'admissions.json').read_text(encoding='utf-8'))

    workbook = openpyxl.load_workbook(workbook_path, read_only=True, data_only=True)
    sheet = workbook['data']
    rows = list(sheet.iter_rows(min_row=2, values_only=True))
    source_keys = {
        (clean(row[3]), clean(row[7]))
        for row in rows
        if clean(row[3]) and clean(row[7]) and clean(row[7]) != '-'
    }
    baseline_keys = {(clean(row.get('u')), clean(row.get('m'))) for row in baseline}
    new_keys = sorted(source_keys - baseline_keys)
    removed_keys = sorted(baseline_keys - source_keys)
    source_signatures = collections.Counter(
        signature(source_record(row))
        for row in rows
        if clean(row[3]) and clean(row[7]) and clean(row[7]) != '-'
    )
    baseline_signatures = collections.Counter(signature(baseline_core(row)) for row in baseline)
    exact_rows = sum((source_signatures & baseline_signatures).values())
    remaining_source = source_signatures.copy()
    unmatched_baseline = []
    for row in baseline:
        row_signature = signature(baseline_core(row))
        if remaining_source[row_signature] > 0:
            remaining_source[row_signature] -= 1
        else:
            unmatched_baseline.append({
                'university': row.get('u', ''),
                'major': row.get('m', ''),
                'officialCorrection': row.get('k1metric', '') == 'donggukOfficial',
            })

    report = {
        'baseline': {
            'version': 'V10.6',
            'rows': len(baseline),
            'uniqueUniversityMajorKeys': len(baseline_keys),
        },
        'workbook': {
            'file': workbook_path.name,
            'sha256': sha256(workbook_path),
            'sheetCount': len(workbook.sheetnames),
            'dataRowsExcludingHeader': len(rows),
            'unique2027UniversityMajorKeys': len(source_keys),
        },
        'guidebook': {
            'file': guidebook_path.name,
            'sha256': sha256(guidebook_path),
            'pages': len(PdfReader(guidebook_path).pages),
        },
        'comparison': {
            'baselineKeysStillPresent': len(baseline_keys & source_keys),
            'baselineKeysMissingFromWorkbook': len(removed_keys),
            'workbookKeysNotInBaseline': len(new_keys),
            'exactBaselineRowsPresent': exact_rows,
            'baselineRowsChangedCorrectedOrUnmatched': len(baseline) - exact_rows,
            'unmatchedRowsWithDonggukOfficialCorrection': sum(
                1 for item in unmatched_baseline if item['officialCorrection']
            ),
            'newKeys': [{'university': item[0], 'major': item[1]} for item in new_keys],
            'removedKeys': [{'university': item[0], 'major': item[1]} for item in removed_keys],
            'unmatchedBaselineRows': unmatched_baseline,
        },
        'decision': 'Do not replace the deployed baseline until the lost mapping and correction rules are reconstructed and validated.'
    }

    output = json.dumps(report, ensure_ascii=False, indent=2) + '\n'
    if args.output:
        output_path = Path(args.output)
        if not output_path.is_absolute():
            output_path = root / output_path
        output_path.parent.mkdir(parents=True, exist_ok=True)
        output_path.write_text(output, encoding='utf-8')
        print(output_path)
    else:
        print(output)


if __name__ == '__main__':
    main()
